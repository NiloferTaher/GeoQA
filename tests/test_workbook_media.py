"""Tests for workbook attachments.

The behaviour worth protecting: the module stays domain-neutral. These tests use
photographs, scans and signatures only as examples of *artefacts*, and assert
nothing that would break if the domain changed.
"""

from __future__ import annotations

import io
import zipfile

import pytest

from geoqa.validations.workbook_media import (
    DEFAULT_MIN_DIMENSION_PX,
    detect_image_format,
    profile_workbook_media,
    validate_workbook_media,
)

openpyxl = pytest.importorskip("openpyxl")


def _names(issues):
    return [i.problem_name for i in issues]


def _png(width=400, height=300):
    """A genuinely decodable PNG, because openpyxl goes through PIL."""
    from io import BytesIO

    from openpyxl.drawing.image import Image
    from PIL import Image as PILImage

    buffer = BytesIO()
    PILImage.new("RGB", (width, height), (200, 200, 200)).save(buffer, format="PNG")
    buffer.seek(0)
    return Image(buffer)


def _book(tmp_path, rows, *, images=None, name="book.xlsx", sheet="Sheet1"):
    from openpyxl import Workbook

    book = Workbook()
    worksheet = book.active
    worksheet.title = sheet
    for record in rows:
        worksheet.append(record)
    for anchor in images or []:
        image, row = anchor
        worksheet.add_image(image, anchor=f"A{row}")
    path = tmp_path / name
    book.save(path)
    return path


# --- format probing ------------------------------------------------------

def test_image_format_detection():
    assert detect_image_format(b"\x89PNG\r\n\x1a\n") == "png"
    assert detect_image_format(b"\xff\xd8\xff\xe0") == "jpeg"
    assert detect_image_format(b"nothing") is None


# --- embedding and anchoring ----------------------------------------------

def test_attachment_is_anchored_to_its_row(tmp_path):
    path = _book(tmp_path, [("ref", "note"), ("R1", "ok")], images=[(_png(), 2)])
    profile = profile_workbook_media(path)
    assert profile.attachment_count == 1
    assert profile.embedded[0].row == 2
    assert profile.attached_rows == (2,)


def test_rows_without_attachment(tmp_path):
    path = _book(tmp_path, [("ref",), ("R1",), ("R2",), ("R3",)], images=[(_png(), 2)])
    profile = profile_workbook_media(path)
    assert profile.rows_without_attachment() == (3, 4)


def test_two_rows_two_attachments(tmp_path):
    path = _book(
        tmp_path,
        [("ref",), ("R1",), ("R2",)],
        images=[(_png(), 2), (_png(), 3)],
    )
    assert profile_workbook_media(path).attachment_count == 2


# --- legibility -----------------------------------------------------------

def test_small_attachment_is_high(tmp_path):
    """Too small to show detail: the row looks evidenced and is not."""
    path = _book(tmp_path, [("ref",), ("R1",)], images=[(_png(40, 30), 2)])
    issues = validate_workbook_media(path, min_dimension_px=DEFAULT_MIN_DIMENSION_PX)
    issue = next(i for i in issues if i.problem_name == "workbook_media_below_legibility_threshold")
    assert str(issue.severity) == "high"


def test_large_attachment_is_accepted(tmp_path):
    path = _book(tmp_path, [("ref",), ("R1",)], images=[(_png(800, 600), 2)])
    assert "workbook_media_below_legibility_threshold" not in _names(validate_workbook_media(path))


def test_legibility_threshold_is_a_parameter(tmp_path):
    path = _book(tmp_path, [("ref",), ("R1",)], images=[(_png(200, 200), 2)])
    assert "workbook_media_below_legibility_threshold" in _names(
        validate_workbook_media(path, min_dimension_px=1000)
    )


# --- broken references -----------------------------------------------------

def test_missing_reference_is_high(tmp_path):
    from openpyxl import Workbook

    book = Workbook()
    worksheet = book.active
    worksheet.append(("ref", "attachment"))
    worksheet.append(("R1", '=IMAGE("C:\\\\photos\\\\R1.jpg")'))
    path = tmp_path / "refs.xlsx"
    book.save(path)
    issues = validate_workbook_media(path, reference_root=tmp_path)
    issue = next(i for i in issues if i.problem_name == "workbook_media_missing_reference")
    assert str(issue.severity) == "high"


def test_existing_reference_is_not_reported(tmp_path):
    from openpyxl import Workbook

    (tmp_path / "R1.jpg").write_bytes(b"\xff\xd8\xff\xe0")
    book = Workbook()
    worksheet = book.active
    worksheet.append(("ref", "attachment"))
    worksheet.append(("R1", '=IMAGE("R1.jpg")'))
    path = tmp_path / "refs.xlsx"
    book.save(path)
    assert "workbook_media_missing_reference" not in _names(
        validate_workbook_media(path, reference_root=tmp_path)
    )


def test_plain_path_cell_is_treated_as_a_reference(tmp_path):
    from openpyxl import Workbook

    book = Workbook()
    worksheet = book.active
    worksheet.append(("ref", "attachment"))
    worksheet.append(("R1", "C:\\photos\\missing.jpg"))
    path = tmp_path / "plainpath.xlsx"
    book.save(path)
    assert "workbook_media_missing_reference" in _names(
        validate_workbook_media(path, reference_root=tmp_path)
    )


# --- attribution -----------------------------------------------------------

def test_attachment_without_key_is_high(tmp_path):
    path = _book(tmp_path, [("ref", "note"), ("R1", "ok"), (None, "ok")], images=[(_png(), 3)])
    issues = validate_workbook_media(path, key_column="ref")
    issue = next(i for i in issues if i.problem_name == "workbook_attachment_without_key")
    assert issue.to_dict()["row"] == 3


def test_attached_row_with_key_is_accepted(tmp_path):
    path = _book(tmp_path, [("ref", "note"), ("R1", "ok")], images=[(_png(), 2)])
    assert "workbook_attachment_without_key" not in _names(
        validate_workbook_media(path, key_column="ref")
    )


def test_missing_key_column_is_reported(tmp_path):
    path = _book(tmp_path, [("ref",), ("R1",)], images=[(_png(), 2)])
    assert "workbook_media_unverifiable" in _names(validate_workbook_media(path, key_column="nope"))


# --- completeness ----------------------------------------------------------

def test_row_missing_attachment_is_opt_in(tmp_path):
    """Plenty of spreadsheets need no attachments, so absence is not a defect."""
    path = _book(tmp_path, [("ref",), ("R1",), ("R2",)], images=[(_png(), 2)])
    assert "workbook_row_missing_attachment" not in _names(validate_workbook_media(path))
    assert "workbook_row_missing_attachment" in _names(
        validate_workbook_media(path, expect_attachment_every_row=True)
    )


def test_complete_sheet_is_quiet(tmp_path):
    path = _book(tmp_path, [("ref",), ("R1",), ("R2",)], images=[(_png(600, 400), 2), (_png(600, 400), 3)])
    issues = validate_workbook_media(
        path, key_column="ref", expect_attachment_every_row=True, sheet="Sheet1"
    )
    assert [i for i in issues if str(i.severity) in {"critical", "high"}] == []


def test_workbook_with_no_media_is_informational(tmp_path):
    path = _book(tmp_path, [("ref",), ("R1",)])
    issue = next(i for i in validate_workbook_media(path) if i.problem_name == "workbook_no_attachments")
    assert str(issue.severity) == "info"


# --- domain neutrality -----------------------------------------------------

def test_module_carries_no_domain_vocabulary():
    """The checks are about artefacts, not about what the artefacts show."""
    import re

    import geoqa.validations.workbook_media as module

    words = set(re.findall(r"[a-z]+", (module.__doc__ or "").lower()))
    for domain_word in ("meter", "water", "network", "leak", "dma", "customer", "utility", "pipe"):
        assert domain_word not in words, domain_word


# --- structure ------------------------------------------------------------

def test_profile_reads_the_package_directly(tmp_path):
    path = _book(tmp_path, [("ref",), ("R1",)], images=[(_png(), 2)])
    with zipfile.ZipFile(path) as archive:
        assert any(n.startswith("xl/media/") for n in archive.namelist())


def test_profile_summary_shape(tmp_path):
    path = _book(tmp_path, [("ref",), ("R1",)], images=[(_png(), 2)])
    summary = profile_workbook_media(path).as_dict()
    assert summary["embedded"] == 1
    assert summary["attachment_rows"] == [2]


def test_specific_sheet_is_honoured(tmp_path):
    path = _book(tmp_path, [("ref",), ("R1",)], images=[(_png(), 2)])
    assert profile_workbook_media(path, sheet="Sheet1").attachment_count == 1