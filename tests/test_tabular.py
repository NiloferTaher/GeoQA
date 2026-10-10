"""Tests for CSV and Excel validation.

The behaviour worth protecting: the failure modes that produce *plausible*
wrong data are caught, and the ones that merely look untidy are not reported.
A table that trips three hundred "issues" gets ignored, which is the same as
having no validator.
"""

from __future__ import annotations

import pytest

from geoqa.validations.tabular import (
    NULL_TOKENS,
    profile_tabular,
    read_csv_rows,
    validate_tabular,
)

openpyxl = pytest.importorskip("openpyxl")


def _names(issues):
    return [i.problem_name for i in issues]


def _csv(tmp_path, text, name="data.csv"):
    path = tmp_path / name
    path.write_text(text, encoding="utf-8")
    return path


def _xlsx(tmp_path, rows, *, name="book.xlsx", sheet="Sheet1"):
    from openpyxl import Workbook

    book = Workbook()
    worksheet = book.active
    worksheet.title = sheet
    for record in rows:
        worksheet.append(record)
    path = tmp_path / name
    book.save(path)
    return path


# --- leading zeros: the headline condition ---------------------------------

def test_leading_zero_identifier_is_critical(tmp_path):
    """007123 read as a number is 7123, which matches nothing at the other end."""
    path = _csv(tmp_path, "account_no,region\n007123,north\n007124,south\n004567,east\n")
    issues = validate_tabular(path)
    issue = next(i for i in issues if i.problem_name == "tabular_identifier_leading_zero")
    assert str(issue.severity) == "critical"
    assert "007123" in issue.to_dict()["samples"]


def test_long_numeric_identifier_is_high(tmp_path):
    path = _csv(tmp_path, "meter_no\n123456789012345678\n987654321098765432\n")
    issues = validate_tabular(path)
    assert "tabular_identifier_numeric" in _names(issues)


def test_ordinary_numeric_column_is_not_an_identifier(tmp_path):
    path = _csv(tmp_path, "area_m2\n1234\n5678\n")
    assert "tabular_identifier_leading_zero" not in _names(validate_tabular(path))
    assert "tabular_identifier_numeric" not in _names(validate_tabular(path))


def test_mixed_type_identifier_is_not_flagged(tmp_path):
    path = _csv(tmp_path, "account_no\n007123\nABC-99\n")
    assert "tabular_identifier_leading_zero" not in _names(validate_tabular(path))


def test_regular_number_in_an_id_column_is_clean(tmp_path):
    path = _csv(tmp_path, "account_no\n7123\n7124\n7125\n")
    assert not [n for n in _names(validate_tabular(path)) if "identifier" in n]


# --- ambiguous dates -------------------------------------------------------

def test_ambiguous_date_is_high(tmp_path):
    """03/04/2025 is 3 April or 4 March. The file does not say which."""
    path = _csv(tmp_path, "read_date\n03/04/2025\n05/06/2025\n")
    issues = validate_tabular(path)
    issue = next(i for i in issues if i.problem_name == "tabular_ambiguous_date")
    assert str(issue.severity) == "high"


def test_unambiguous_date_is_clean(tmp_path):
    path = _csv(tmp_path, "read_date\n2025-04-03\n2025-06-05\n")
    assert "tabular_ambiguous_date" not in _names(validate_tabular(path))


def test_ymd_date_is_clean(tmp_path):
    path = _csv(tmp_path, "read_date\n2025/04/03\n")
    assert "tabular_ambiguous_date" not in _names(validate_tabular(path))


# --- sentinel values -------------------------------------------------------

def test_sentinel_values_are_flagged(tmp_path):
    """-9999 is an ordinary number and will enter a mean unless excluded."""
    path = _csv(tmp_path, "pressure_bar\n4.2\n-9999\n4.5\n")
    issues = validate_tabular(path)
    assert "tabular_sentinel_values" in _names(issues)


def test_normal_negatives_are_not_sentinels(tmp_path):
    path = _csv(tmp_path, "elevation_m\n120.5\n-15.2\n")
    assert "tabular_sentinel_values" not in _names(validate_tabular(path))


# --- structure --------------------------------------------------------------

def test_duplicate_header(tmp_path):
    path = _csv(tmp_path, "name,region,region\nA,north,south\n")
    assert "tabular_duplicate_header" in _names(validate_tabular(path))


def test_blank_header(tmp_path):
    path = _csv(tmp_path, "name,,region\nA,x,north\n")
    assert "tabular_blank_header" in _names(validate_tabular(path))


def test_ragged_rows(tmp_path):
    path = _csv(tmp_path, "a,b,c\n1,2,3\n4,5\n6,7,8,9\n")
    issues = validate_tabular(path)
    issue = next(i for i in issues if i.problem_name == "tabular_ragged_rows")
    assert issue.to_dict()["rows_affected"] == 2


def test_missing_expected_column(tmp_path):
    path = _csv(tmp_path, "a,b\n1,2\n")
    issue = next(i for i in validate_tabular(path, expected_columns=["c"]) if i.problem_name == "tabular_missing_expected_column")
    assert issue.to_dict()["column"] == "c"


def test_empty_column(tmp_path):
    path = _csv(tmp_path, "a,b\n1,\n2,\n")
    assert "tabular_empty_column" in _names(validate_tabular(path))


def test_missing_key_is_critical(tmp_path):
    path = _csv(tmp_path, "a\n1\n")
    issues = validate_tabular(path, key_column="id", require_key=True)
    issue = next(i for i in issues if i.problem_name == "tabular_missing_key")
    assert str(issue.severity) == "critical"


def test_duplicate_key_is_high(tmp_path):
    path = _csv(tmp_path, "id,region\nA,north\nA,south\nB,east\n")
    issues = validate_tabular(path, key_column="id")
    issue = next(i for i in issues if i.problem_name == "tabular_duplicate_key")
    assert issue.to_dict()["unique"] == 2


def test_unique_key_is_clean(tmp_path):
    path = _csv(tmp_path, "id,region\nA,north\nB,south\n")
    assert "tabular_duplicate_key" not in _names(validate_tabular(path, key_column="id"))


# --- Excel-specific ----------------------------------------------------------

def test_merged_cells_are_high(tmp_path):
    """A merged cell stores its value only on the first row."""
    from openpyxl import Workbook

    book = Workbook()
    worksheet = book.active
    for record in [("region", "value"), ("north", 1), ("south", 2)]:
        worksheet.append(record)
    worksheet.merge_cells("A2:A3")
    path = tmp_path / "merged.xlsx"
    book.save(path)
    issues = validate_tabular(path)
    assert "tabular_merged_cells" in _names(issues)


def test_hidden_sheet_is_high(tmp_path):
    from openpyxl import Workbook

    book = Workbook()
    book.active.title = "Visible"
    book.create_sheet("Corrections")
    book["Corrections"].sheet_state = "hidden"
    book["Corrections"].append(("id", "value"))
    book["Corrections"].append(("A", 1))
    path = tmp_path / "hidden.xlsx"
    book.save(path)
    issue = next(i for i in validate_tabular(path) if i.problem_name == "tabular_hidden_sheet")
    assert issue.to_dict()["sheet"] == "Corrections"


def test_hidden_rows_are_reported(tmp_path):
    from openpyxl import Workbook

    book = Workbook()
    worksheet = book.active
    for record in [("a", "b"), ("1", "2"), ("3", "4")]:
        worksheet.append(record)
    worksheet.row_dimensions[3].hidden = True
    path = tmp_path / "hiddenrows.xlsx"
    book.save(path)
    assert "tabular_hidden_rows" in _names(validate_tabular(path))


def test_formula_cells_are_reported(tmp_path):
    from openpyxl import Workbook

    book = Workbook()
    worksheet = book.active
    worksheet.append(("a", "b"))
    worksheet.append((1, "=A2*2"))
    path = tmp_path / "formula.xlsx"
    book.save(path)
    assert "tabular_formula_cells" in _names(validate_tabular(path))


def test_clean_workbook_is_quiet(tmp_path):
    path = _xlsx(tmp_path, [("id", "region"), ("A", "north"), ("B", "south")])
    serious = [i for i in validate_tabular(path) if str(i.severity) in {"critical", "high"}]
    assert serious == []


def test_multiple_sheets_are_listed(tmp_path):
    from openpyxl import Workbook

    book = Workbook()
    book.active.title = "One"
    book.create_sheet("Two")
    path = tmp_path / "multi.xlsx"
    book.save(path)
    assert profile_tabular(path).sheets == ("One", "Two")


def test_specific_sheet_can_be_validated(tmp_path):
    from openpyxl import Workbook

    book = Workbook()
    book.active.title = "One"
    book.active.append(("a",))
    second = book.create_sheet("Two")
    second.append(("id", "region"))
    second.append(("007123", "north"))
    path = tmp_path / "multi.xlsx"
    book.save(path)
    assert "tabular_identifier_leading_zero" in _names(validate_tabular(path, sheet="Two"))
    assert "tabular_identifier_leading_zero" not in _names(validate_tabular(path, sheet="One"))


# --- CSV dialect ---------------------------------------------------------------

def test_semicolon_delimiter_is_detected(tmp_path):
    path = _csv(tmp_path, "id;region\nA;north\nB;south\n")
    assert profile_tabular(path).delimiter == ";"


def test_tab_delimiter_is_detected(tmp_path):
    path = _csv(tmp_path, "id\tregion\nA\tnorth\n")
    assert profile_tabular(path).delimiter == "\t"


def test_utf8_bom_is_handled(tmp_path):
    path = tmp_path / "bom.csv"
    path.write_bytes("id,region\nA,north\n".encode("utf-8-sig"))
    assert profile_tabular(path).column_names == ("id", "region")


def test_null_tokens_are_recognised(tmp_path):
    path = _csv(tmp_path, "a,b\n1,N/A\n2,-\n3,NULL\n")
    profile = profile_tabular(path)
    assert profile.columns[1].null_count == 3


def test_null_token_vocabulary_is_deliberate():
    assert "n/a" in NULL_TOKENS
    assert "null" in NULL_TOKENS


def test_unsupported_extension_raises(tmp_path):
    path = tmp_path / "data.dbf"
    path.write_bytes(b"\x00")
    with pytest.raises(ValueError, match="not a recognised tabular file"):
        profile_tabular(path)


# --- reading stays textual -----------------------------------------------------

def test_read_csv_rows_keeps_strings(tmp_path):
    """Coercion is where identifiers die, so it has to be asked for."""
    path = _csv(tmp_path, "account_no,region\n007123,north\n")
    rows = read_csv_rows(path)
    assert rows[0]["account_no"] == "007123"


def test_read_csv_rows_respects_limit(tmp_path):
    path = _csv(tmp_path, "a\n1\n2\n3\n4\n")
    assert len(read_csv_rows(path, max_rows=2)) == 2


def test_profile_row_count(tmp_path):
    path = _csv(tmp_path, "a,b\n1,2\n3,4\n5,6\n")
    assert profile_tabular(path).rows == 3


def test_blank_lines_are_skipped(tmp_path):
    path = _csv(tmp_path, "a,b\n1,2\n\n3,4\n")
    assert profile_tabular(path).rows == 2